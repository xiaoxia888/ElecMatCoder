#!/usr/bin/env python3
"""Analyze routing quality, confidence calibration, and code correctness."""

from __future__ import annotations

import argparse
import logging
import math
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd


LOGGER = logging.getLogger("routing_confidence_analyzer")

PREDICTED_COLUMN = "excel2_原始总编码"
TRUTH_COLUMN = "本项目材料代码"
CONFIDENCE_CANDIDATES = ("excel2_模型置信分", "excel2_总置信度")
DIFFICULTY_PREFIX = "excel2_分流最终难度"
PROJECT_CANDIDATES = ("excel2_项目名称", "项目简称")
REASON_COLUMN = "excel2_分流原因"
REVIEW_COLUMN = "excel2_是否需审核"

FIELD_COLUMNS = {
    "TYPE": ("C1-名称简写", "excel2_TYPE_原始编码"),
    "SIZE": ("C1-规格简写", "excel2_SIZE_原始编码"),
    "THICKNESS": ("C1-壁厚简写", "excel2_THICKNESS_原始编码"),
    "PRESSURE": ("C1-压力等级简写", "excel2_PRESSURE_原始编码"),
    "MATERIAL": ("C1-材质简写", "excel2_MATERIAL_原始编码"),
    "STANDARD": ("C1-标准简写", "excel2_STANDARD_原始编码"),
}

FIXED_CONFIDENCE_BINS = [-0.000001, 0.5, 0.8, 0.9, 0.95, 0.98, 0.99, 0.995, 0.999, 1.000001]
FIXED_CONFIDENCE_LABELS = [
    "0-0.50",
    "0.50-0.80",
    "0.80-0.90",
    "0.90-0.95",
    "0.95-0.98",
    "0.98-0.99",
    "0.99-0.995",
    "0.995-0.999",
    "0.999-1.0",
]


@dataclass(frozen=True)
class AnalysisConfig:
    input_path: Path
    output_dir: Path
    sheet: str | int
    project_column: str | None
    reference_threshold: float
    target_accuracies: tuple[float, ...]
    min_project_samples: int
    overwrite: bool


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="分析分流、模型置信分与编码正确率的交集关系")
    parser.add_argument("--input", required=True, type=Path, help="运行结果 Excel 路径")
    parser.add_argument("--output-dir", required=True, type=Path, help="Excel 和分析图片输出目录")
    parser.add_argument("--sheet", default="0", help="工作表名称或从 0 开始的序号")
    parser.add_argument("--project-column", help="项目列名；默认优先使用 excel2_项目名称")
    parser.add_argument("--reference-threshold", type=float, default=0.99, help="四象限参考置信分阈值")
    parser.add_argument(
        "--target-accuracy",
        type=float,
        action="append",
        dest="target_accuracies",
        help="自动放行目标准确率，可重复；默认 0.99、0.995、0.999",
    )
    parser.add_argument("--min-project-samples", type=int, default=30, help="项目异常判定最少样本数")
    parser.add_argument("--overwrite", action="store_true", help="覆盖输出目录中的同名结果")
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return parser.parse_args()


def build_config(args: argparse.Namespace) -> AnalysisConfig:
    sheet: str | int = args.sheet
    if isinstance(sheet, str) and sheet.isdigit():
        sheet = int(sheet)
    targets = tuple(args.target_accuracies or (0.99, 0.995, 0.999))
    for value in (*targets, args.reference_threshold):
        if not 0 <= value <= 1:
            raise ValueError("准确率和置信分阈值必须在 0 到 1 之间")
    if args.min_project_samples <= 0:
        raise ValueError("--min-project-samples 必须大于 0")
    return AnalysisConfig(
        input_path=args.input.expanduser().resolve(),
        output_dir=args.output_dir.expanduser().resolve(),
        sheet=sheet,
        project_column=args.project_column,
        reference_threshold=float(args.reference_threshold),
        target_accuracies=targets,
        min_project_samples=int(args.min_project_samples),
        overwrite=bool(args.overwrite),
    )


def text(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    if isinstance(value, (float, np.floating)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_code(value: Any) -> str:
    normalized = unicodedata.normalize("NFKC", text(value)).upper()
    return re.sub(r"\s+", "", normalized)


def safe_rate(numerator: int | float, denominator: int | float) -> float:
    return float(numerator / denominator) if denominator else math.nan


def resolve_column(columns: Iterable[str], candidates: Iterable[str], *, prefix: str | None = None) -> str:
    available = list(columns)
    for candidate in candidates:
        if candidate in available:
            return candidate
    if prefix:
        matches = [column for column in available if str(column).startswith(prefix)]
        if matches:
            return matches[0]
    expected = ", ".join(candidates)
    raise ValueError(f"Excel 缺少必要列: {expected or prefix}")


def parse_confidence(series: pd.Series) -> pd.Series:
    raw = series.map(text)
    percent = raw.str.endswith("%")
    values = pd.to_numeric(raw.str.rstrip("%"), errors="coerce")
    values.loc[percent] = values.loc[percent] / 100.0
    values.loc[values > 1.5] = values.loc[values > 1.5] / 100.0
    return values.clip(0.0, 1.0)


def load_and_prepare(config: AnalysisConfig) -> tuple[pd.DataFrame, dict[str, str], pd.DataFrame]:
    if not config.input_path.is_file():
        raise FileNotFoundError(f"输入 Excel 不存在: {config.input_path}")
    source = pd.read_excel(config.input_path, sheet_name=config.sheet, dtype=object)
    if source.empty:
        raise ValueError("输入 Excel 没有数据")
    required = {PREDICTED_COLUMN, TRUTH_COLUMN, REASON_COLUMN}
    missing = required - set(source.columns)
    if missing:
        raise ValueError(f"Excel 缺少必要列: {sorted(missing)}")

    confidence_column = resolve_column(source.columns, CONFIDENCE_CANDIDATES)
    difficulty_column = resolve_column(source.columns, (), prefix=DIFFICULTY_PREFIX)
    project_column = config.project_column or resolve_column(source.columns, PROJECT_CANDIDATES)
    if project_column not in source.columns:
        raise ValueError(f"Excel 缺少项目列: {project_column}")

    detail = source.copy()
    detail["分析_Excel行号"] = np.arange(2, len(detail) + 2)
    detail["分析_预测编码标准化"] = detail[PREDICTED_COLUMN].map(normalize_code)
    detail["分析_正确编码标准化"] = detail[TRUTH_COLUMN].map(normalize_code)
    detail["分析_编码是否正确"] = detail["分析_预测编码标准化"].eq(detail["分析_正确编码标准化"])
    detail["分析_模型置信分"] = parse_confidence(detail[confidence_column])
    difficulty_number = pd.to_numeric(detail[difficulty_column], errors="coerce")
    detail["分析_分流数值"] = difficulty_number
    detail["分析_分流类别"] = difficulty_number.map({0.0: "困难", 1.0: "中等", 2.0: "简单"}).fillna("未知")
    detail["分析_项目名称"] = detail[project_column].map(text).replace("", "（项目为空）")
    detail["分析_分流原因"] = detail[REASON_COLUMN].map(text).replace("", "（原因为空）")
    detail["分析_参考高置信"] = detail["分析_模型置信分"].ge(config.reference_threshold)
    detail["分析_四象限"] = (
        detail["分析_分流类别"]
        + " + "
        + detail["分析_参考高置信"].map({True: "高置信", False: "低置信"})
    )

    quality_rows = []
    for column in (PREDICTED_COLUMN, TRUTH_COLUMN, confidence_column, difficulty_column, project_column, REASON_COLUMN):
        blank = source[column].map(text).eq("")
        quality_rows.append({"字段": column, "缺失或空值数": int(blank.sum()), "缺失率": float(blank.mean())})
    quality = pd.DataFrame(quality_rows)
    columns = {
        "confidence": confidence_column,
        "difficulty": difficulty_column,
        "project": project_column,
    }
    return detail, columns, quality


def subset_metrics(frame: pd.DataFrame, total_errors: int) -> dict[str, Any]:
    count = len(frame)
    correct = int(frame["分析_编码是否正确"].sum())
    errors = count - correct
    return {
        "样本数": count,
        "样本占比": math.nan,
        "正确数": correct,
        "错误数": errors,
        "准确率": safe_rate(correct, count),
        "错误率": safe_rate(errors, count),
        "全部错误捕获占比": safe_rate(errors, total_errors),
        "平均置信分": float(frame["分析_模型置信分"].mean()) if count else math.nan,
        "置信分中位数": float(frame["分析_模型置信分"].median()) if count else math.nan,
    }


def build_routing_stats(detail: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    total = len(detail)
    total_correct = int(detail["分析_编码是否正确"].sum())
    total_errors = total - total_correct
    rows = []
    for route in ("简单", "困难", "中等", "未知"):
        group = detail.loc[detail["分析_分流类别"].eq(route)]
        if group.empty:
            continue
        metrics = subset_metrics(group, total_errors)
        metrics["分流类别"] = route
        metrics["样本占比"] = safe_rate(len(group), total)
        rows.append(metrics)
    routing = pd.DataFrame(rows)[
        ["分流类别", "样本数", "样本占比", "正确数", "错误数", "准确率", "错误率", "全部错误捕获占比", "平均置信分", "置信分中位数"]
    ]

    simple = detail.loc[detail["分析_分流类别"].eq("简单")]
    hard = detail.loc[detail["分析_分流类别"].eq("困难")]
    hard_errors = int((~hard["分析_编码是否正确"]).sum())
    summary = pd.DataFrame([
        {"指标": "总样本数", "值": total, "说明": "参与分析的全部行"},
        {"指标": "项目数", "值": detail["分析_项目名称"].nunique(), "说明": "按项目名称去重"},
        {"指标": "总体准确率", "值": safe_rate(total_correct, total), "说明": "预测总编码与项目正确编码直接比对"},
        {"指标": "总体错误数", "值": total_errors, "说明": "预测总编码不一致的行数"},
        {"指标": "简单覆盖率", "值": safe_rate(len(simple), total), "说明": "当前可自动放行的数据占比"},
        {"指标": "简单准确率", "值": float(simple["分析_编码是否正确"].mean()), "说明": "当前自动放行样本准确率"},
        {"指标": "简单错误泄漏数", "值": int((~simple["分析_编码是否正确"]).sum()), "说明": "被标简单但编码错误"},
        {"指标": "困难准确率", "值": float(hard["分析_编码是否正确"].mean()), "说明": "当前困难样本准确率"},
        {"指标": "困难错误捕获率", "值": safe_rate(hard_errors, total_errors), "说明": "全部错误中被困难分流拦截的比例"},
        {"指标": "置信分均值", "值": float(detail["分析_模型置信分"].mean()), "说明": "原模型置信分均值"},
    ])
    return summary, routing


def _group_accuracy(frame: pd.DataFrame) -> tuple[int, int, float]:
    count = len(frame)
    correct = int(frame["分析_编码是否正确"].sum())
    return count, count - correct, safe_rate(correct, count)


def build_project_stats(detail: pd.DataFrame, min_samples: int) -> pd.DataFrame:
    global_accuracy = float(detail["分析_编码是否正确"].mean())
    global_simple = detail.loc[detail["分析_分流类别"].eq("简单")]
    global_hard = detail.loc[detail["分析_分流类别"].eq("困难")]
    global_simple_share = safe_rate(len(global_simple), len(detail))
    global_hard_share = safe_rate(len(global_hard), len(detail))
    global_simple_accuracy = float(global_simple["分析_编码是否正确"].mean())
    global_hard_accuracy = float(global_hard["分析_编码是否正确"].mean())
    rows = []
    for project, group in detail.groupby("分析_项目名称", dropna=False, sort=False):
        simple = group.loc[group["分析_分流类别"].eq("简单")]
        hard = group.loc[group["分析_分流类别"].eq("困难")]
        count, errors, accuracy = _group_accuracy(group)
        simple_count, simple_errors, simple_accuracy = _group_accuracy(simple)
        hard_count, hard_errors, hard_accuracy = _group_accuracy(hard)
        simple_share = safe_rate(simple_count, count)
        hard_share = safe_rate(hard_count, count)
        flags: list[str] = []
        if count >= min_samples:
            if hard_share > 0.5:
                flags.append("困难占比超过简单")
            if simple_share >= max(0.75, global_simple_share + 0.10) and simple_accuracy <= min(0.97, global_simple_accuracy - 0.02):
                flags.append("简单占比高但简单准确率低")
            if hard_share >= max(0.40, global_hard_share + 0.10) and hard_accuracy >= max(0.95, global_hard_accuracy + 0.08):
                flags.append("困难占比高但困难准确率高（分流偏保守）")
            if accuracy <= global_accuracy - 0.05:
                flags.append("项目总体准确率显著偏低")
        else:
            flags.append("样本不足，仅观察")
        rows.append({
            "项目名称": project,
            "样本数": count,
            "正确数": count - errors,
            "错误数": errors,
            "总体准确率": accuracy,
            "简单数": simple_count,
            "简单占比": simple_share,
            "简单错误数": simple_errors,
            "简单准确率": simple_accuracy,
            "困难数": hard_count,
            "困难占比": hard_share,
            "困难错误数": hard_errors,
            "困难准确率": hard_accuracy,
            "平均置信分": float(group["分析_模型置信分"].mean()),
            "正确样本平均置信分": float(group.loc[group["分析_编码是否正确"], "分析_模型置信分"].mean()),
            "错误样本平均置信分": float(group.loc[~group["分析_编码是否正确"], "分析_模型置信分"].mean()),
            "异常类型": "；".join(flags),
            "是否重点异常": bool(flags and flags != ["样本不足，仅观察"]),
        })
    return pd.DataFrame(rows).sort_values(["是否重点异常", "错误数", "样本数"], ascending=[False, False, False])


def build_confidence_bins(detail: pd.DataFrame) -> pd.DataFrame:
    work = detail.dropna(subset=["分析_模型置信分"]).copy()
    work["置信分区间"] = pd.cut(
        work["分析_模型置信分"],
        bins=FIXED_CONFIDENCE_BINS,
        labels=FIXED_CONFIDENCE_LABELS,
        include_lowest=True,
        ordered=True,
    )
    rows = []
    for label, group in work.groupby("置信分区间", observed=False):
        count, errors, accuracy = _group_accuracy(group)
        rows.append({
            "置信分区间": str(label),
            "样本数": count,
            "样本占比": safe_rate(count, len(work)),
            "正确数": count - errors,
            "错误数": errors,
            "准确率": accuracy,
            "错误率": safe_rate(errors, count),
            "平均置信分": float(group["分析_模型置信分"].mean()) if count else math.nan,
        })
    return pd.DataFrame(rows)


def roc_auc_score_binary(labels: pd.Series, scores: pd.Series) -> float:
    valid = labels.notna() & scores.notna()
    y = labels.loc[valid].astype(int)
    score = scores.loc[valid].astype(float)
    positives = int(y.sum())
    negatives = len(y) - positives
    if positives == 0 or negatives == 0:
        return math.nan
    ranks = score.rank(method="average")
    return float((ranks.loc[y.eq(1)].sum() - positives * (positives + 1) / 2) / (positives * negatives))


def build_calibration(detail: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    work = detail.dropna(subset=["分析_模型置信分"]).copy()
    bin_count = min(10, max(2, len(work)))
    work["校准分组"] = pd.qcut(
        work["分析_模型置信分"],
        q=bin_count,
        duplicates="drop",
    )
    rows = []
    for group_id, (_, group) in enumerate(work.groupby("校准分组", observed=True, sort=True), start=1):
        accuracy = float(group["分析_编码是否正确"].mean())
        mean_confidence = float(group["分析_模型置信分"].mean())
        rows.append({
            "分组": group_id,
            "样本数": len(group),
            "最低置信分": float(group["分析_模型置信分"].min()),
            "最高置信分": float(group["分析_模型置信分"].max()),
            "平均置信分": mean_confidence,
            "实际准确率": accuracy,
            "校准差（置信分-准确率）": mean_confidence - accuracy,
            "绝对校准差": abs(mean_confidence - accuracy),
        })
    calibration = pd.DataFrame(rows)
    ece = float((calibration["绝对校准差"] * calibration["样本数"] / calibration["样本数"].sum()).sum())
    labels = work["分析_编码是否正确"].astype(float)
    scores = work["分析_模型置信分"].astype(float)
    metrics = pd.DataFrame([
        {"指标": "有效置信分样本数", "值": len(work), "说明": "置信分可解析的样本"},
        {"指标": "正确性区分AUC", "值": roc_auc_score_binary(labels, scores), "说明": "越接近1，置信分越能区分正确与错误"},
        {"指标": "Brier分数", "值": float(np.mean((scores - labels) ** 2)), "说明": "越接近0，概率质量越好"},
        {"指标": "期望校准误差ECE", "值": ece, "说明": "越接近0，置信分越接近实际正确率"},
        {"指标": "正确样本平均置信分", "值": float(work.loc[work["分析_编码是否正确"], "分析_模型置信分"].mean()), "说明": "正确样本分数均值"},
        {"指标": "错误样本平均置信分", "值": float(work.loc[~work["分析_编码是否正确"], "分析_模型置信分"].mean()), "说明": "错误样本分数均值"},
    ])
    return calibration, metrics


def acceptance_metrics(detail: pd.DataFrame, mask: pd.Series) -> dict[str, Any]:
    accepted = detail.loc[mask]
    reviewed = detail.loc[~mask]
    total_errors = int((~detail["分析_编码是否正确"]).sum())
    accepted_errors = int((~accepted["分析_编码是否正确"]).sum())
    reviewed_errors = int((~reviewed["分析_编码是否正确"]).sum())
    return {
        "自动放行数": len(accepted),
        "自动放行覆盖率": safe_rate(len(accepted), len(detail)),
        "自动放行正确数": len(accepted) - accepted_errors,
        "自动放行错误数": accepted_errors,
        "自动放行准确率": safe_rate(len(accepted) - accepted_errors, len(accepted)),
        "审核数": len(reviewed),
        "审核占比": safe_rate(len(reviewed), len(detail)),
        "审核捕获错误数": reviewed_errors,
        "错误拦截率": safe_rate(reviewed_errors, total_errors),
    }


def build_threshold_simulation(detail: pd.DataFrame) -> pd.DataFrame:
    scores = detail["分析_模型置信分"]
    simple = detail["分析_分流类别"].eq("简单")
    rows = []
    route_metrics = acceptance_metrics(detail, simple)
    rows.append({"策略": "仅分流（简单放行）", "置信分阈值": math.nan, **route_metrics})
    for threshold in np.linspace(0.0, 1.0, 1001):
        high = scores.ge(float(threshold)) & scores.notna()
        rows.append({"策略": "仅置信分", "置信分阈值": float(threshold), **acceptance_metrics(detail, high)})
        rows.append({"策略": "分流+置信分交集", "置信分阈值": float(threshold), **acceptance_metrics(detail, simple & high)})
    return pd.DataFrame(rows)


def build_recommendations(simulation: pd.DataFrame, targets: tuple[float, ...]) -> pd.DataFrame:
    rows = []
    for strategy in ("仅分流（简单放行）", "仅置信分", "分流+置信分交集"):
        strategy_rows = simulation.loc[simulation["策略"].eq(strategy)]
        for target in targets:
            eligible = strategy_rows.loc[
                strategy_rows["自动放行准确率"].ge(target) & strategy_rows["自动放行数"].gt(0)
            ].sort_values(["自动放行覆盖率", "自动放行准确率"], ascending=[False, False])
            if eligible.empty:
                rows.append({
                    "策略": strategy,
                    "目标自动放行准确率": target,
                    "是否可达到": False,
                    "推荐置信分阈值": math.nan,
                    "自动放行数": 0,
                    "自动放行覆盖率": 0.0,
                    "实际自动放行准确率": math.nan,
                    "自动放行错误数": math.nan,
                    "错误拦截率": math.nan,
                })
                continue
            best = eligible.iloc[0]
            rows.append({
                "策略": strategy,
                "目标自动放行准确率": target,
                "是否可达到": True,
                "推荐置信分阈值": best["置信分阈值"],
                "自动放行数": int(best["自动放行数"]),
                "自动放行覆盖率": float(best["自动放行覆盖率"]),
                "实际自动放行准确率": float(best["自动放行准确率"]),
                "自动放行错误数": int(best["自动放行错误数"]),
                "错误拦截率": float(best["错误拦截率"]),
            })
    return pd.DataFrame(rows)


def build_quadrants(detail: pd.DataFrame, threshold: float) -> pd.DataFrame:
    rows = []
    for route in ("简单", "困难", "中等", "未知"):
        route_mask = detail["分析_分流类别"].eq(route)
        if not route_mask.any():
            continue
        for confidence_label, high in (("高置信", True), ("低置信", False)):
            group = detail.loc[route_mask & detail["分析_模型置信分"].ge(threshold).eq(high)]
            count, errors, accuracy = _group_accuracy(group)
            rows.append({
                "分流类别": route,
                "置信分组": confidence_label,
                "参考阈值": threshold,
                "样本数": count,
                "样本占比": safe_rate(count, len(detail)),
                "正确数": count - errors,
                "错误数": errors,
                "准确率": accuracy,
                "平均置信分": float(group["分析_模型置信分"].mean()) if count else math.nan,
            })
    return pd.DataFrame(rows)


def build_strategy_comparison(detail: pd.DataFrame, threshold: float) -> pd.DataFrame:
    score_is_high = detail["分析_模型置信分"].ge(threshold) & detail["分析_模型置信分"].notna()
    is_simple = detail["分析_分流类别"].eq("简单")
    definitions = (
        ("仅分流", "简单样本自动放行", is_simple),
        ("仅模型分", f"置信分 >= {threshold:.4f} 自动放行", score_is_high),
        ("分流+模型分交集", f"简单且置信分 >= {threshold:.4f} 才自动放行", is_simple & score_is_high),
    )
    rows = []
    for strategy, rule, mask in definitions:
        rows.append({
            "策略": strategy,
            "自动放行规则": rule,
            "参考置信分阈值": threshold if strategy != "仅分流" else math.nan,
            **acceptance_metrics(detail, mask),
        })
    return pd.DataFrame(rows)


def build_conclusions(
    summary: pd.DataFrame,
    strategy: pd.DataFrame,
    recommendations: pd.DataFrame,
    confidence_metrics: pd.DataFrame,
) -> pd.DataFrame:
    summary_values = dict(zip(summary["指标"], summary["值"]))
    strategy_rows = strategy.set_index("策略")
    route = strategy_rows.loc["仅分流"]
    score = strategy_rows.loc["仅模型分"]
    intersection = strategy_rows.loc["分流+模型分交集"]
    metric_values = dict(zip(confidence_metrics["指标"], confidence_metrics["值"]))
    rows = [
        {
            "序号": 1,
            "主题": "总体模型表现",
            "结论": f"总样本 {int(summary_values['总样本数']):,} 条，总体准确率 {summary_values['总体准确率']:.2%}，错误 {int(summary_values['总体错误数']):,} 条。",
        },
        {
            "序号": 2,
            "主题": "当前仅分流",
            "结论": f"简单样本放行覆盖率 {route['自动放行覆盖率']:.2%}，放行准确率 {route['自动放行准确率']:.2%}，仍会放行 {int(route['自动放行错误数']):,} 条错误。",
        },
        {
            "序号": 3,
            "主题": "仅模型分",
            "结论": f"按参考阈值单独使用模型分，覆盖率 {score['自动放行覆盖率']:.2%}，准确率 {score['自动放行准确率']:.2%}，放行错误 {int(score['自动放行错误数']):,} 条。",
        },
        {
            "序号": 4,
            "主题": "交集策略",
            "结论": f"简单且高置信才放行时，覆盖率 {intersection['自动放行覆盖率']:.2%}，准确率 {intersection['自动放行准确率']:.2%}，放行错误降至 {int(intersection['自动放行错误数']):,} 条。",
        },
        {
            "序号": 5,
            "主题": "三种策略判断",
            "结论": "交集策略不会改变模型总体准确率，但能减少自动放行错误；代价是降低自动放行覆盖率、增加审核量。",
        },
        {
            "序号": 6,
            "主题": "模型分有效性",
            "结论": f"置信分区分正确/错误的 AUC 为 {metric_values.get('正确性区分AUC', math.nan):.4f}，ECE 为 {metric_values.get('期望校准误差ECE', math.nan):.4f}；应以经验阈值使用，不能直接把原始分数当真实正确概率。",
        },
    ]
    target_rows = recommendations.loc[
        recommendations["策略"].eq("分流+置信分交集") & recommendations["是否可达到"].astype(bool)
    ]
    for _, row in target_rows.iterrows():
        rows.append({
            "序号": len(rows) + 1,
            "主题": f"交集策略达到 {row['目标自动放行准确率']:.2%}",
            "结论": f"推荐阈值 {row['推荐置信分阈值']:.4f}，覆盖率 {row['自动放行覆盖率']:.2%}，实际准确率 {row['实际自动放行准确率']:.2%}，放行错误 {int(row['自动放行错误数']):,} 条。",
        })
    return pd.DataFrame(rows)


def build_image_guide(threshold: float) -> pd.DataFrame:
    rows = [
        ("01_三种策略覆盖率准确率对比.png", f"在参考阈值 {threshold:.4f} 下比较仅分流、仅模型分和二者交集。", "覆盖率表示可自动处理多少；准确率表示自动放行有多可靠。"),
        ("02_三种策略错误控制对比.png", "左图看自动放错多少条，右图看能拦住全部错误的多少。", "放行错误越少、错误拦截率越高越安全。"),
        ("03_置信分区间准确率.png", "柱形是每个分数区间的样本数，红线是该区间真实准确率。", "分数越高，真实准确率应总体升高，否则模型分不能可靠排序风险。"),
        ("04_项目分流异常散点.png", "横轴是项目简单占比，纵轴是项目简单样本准确率。", "右下区域最危险：自动放行比例高，但放行准确率低。"),
    ]
    return pd.DataFrame(rows, columns=["图片文件", "怎么看", "判断标准"])


def _reason_table(records: list[dict[str, Any]], total: int) -> pd.DataFrame:
    frame = pd.DataFrame(records)
    rows = []
    for reason, group in frame.groupby("原因", sort=False):
        count = len(group)
        correct = int(group["正确"].sum())
        rows.append({
            "分流原因": reason,
            "样本数": count,
            "样本占比": safe_rate(count, total),
            "正确数": correct,
            "错误数": count - correct,
            "准确率": safe_rate(correct, count),
            "错误率": safe_rate(count - correct, count),
            "平均置信分": float(group["置信分"].mean()),
        })
    return pd.DataFrame(rows).sort_values(["错误数", "样本数"], ascending=False)


def build_reason_stats(detail: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    combined_records = [
        {"原因": row["分析_分流原因"], "正确": row["分析_编码是否正确"], "置信分": row["分析_模型置信分"]}
        for _, row in detail.iterrows()
    ]
    atomic_records = []
    for _, row in detail.iterrows():
        reasons = [part.strip() for part in str(row["分析_分流原因"]).split("|") if part.strip()]
        for reason in reasons or ["（原因为空）"]:
            atomic_records.append({"原因": reason, "正确": row["分析_编码是否正确"], "置信分": row["分析_模型置信分"]})
    return _reason_table(combined_records, len(detail)), _reason_table(atomic_records, len(detail))


def build_field_stats(detail: pd.DataFrame) -> pd.DataFrame:
    total_errors = int((~detail["分析_编码是否正确"]).sum())
    rows = []
    for field, (truth_column, predicted_column) in FIELD_COLUMNS.items():
        if truth_column not in detail.columns or predicted_column not in detail.columns:
            continue
        truth = detail[truth_column].map(normalize_code)
        predicted = detail[predicted_column].map(normalize_code)
        relevant = truth.ne("") | predicted.ne("")
        matches = truth.eq(predicted)
        mismatch = relevant & ~matches
        mismatch_in_wrong = mismatch & ~detail["分析_编码是否正确"]
        rows.append({
            "字段": field,
            "正确编码列": truth_column,
            "模型编码列": predicted_column,
            "有效样本数": int(relevant.sum()),
            "字段一致数": int((relevant & matches).sum()),
            "字段不一致数": int(mismatch.sum()),
            "字段准确率": safe_rate(int((relevant & matches).sum()), int(relevant.sum())),
            "总编码错误中的字段不一致数": int(mismatch_in_wrong.sum()),
            "对总错误的覆盖占比": safe_rate(int(mismatch_in_wrong.sum()), total_errors),
        })
    if not rows:
        return pd.DataFrame(columns=[
            "字段",
            "正确编码列",
            "模型编码列",
            "有效样本数",
            "字段一致数",
            "字段不一致数",
            "字段准确率",
            "总编码错误中的字段不一致数",
            "对总错误的覆盖占比",
        ])
    return pd.DataFrame(rows).sort_values("字段不一致数", ascending=False)


def build_review_cross(detail: pd.DataFrame) -> pd.DataFrame:
    if REVIEW_COLUMN not in detail.columns:
        return pd.DataFrame()
    rows = []
    for (route, review), group in detail.groupby(["分析_分流类别", REVIEW_COLUMN], dropna=False):
        count, errors, accuracy = _group_accuracy(group)
        rows.append({
            "分流类别": route,
            "是否需审核": text(review),
            "样本数": count,
            "错误数": errors,
            "准确率": accuracy,
        })
    return pd.DataFrame(rows)


def configure_matplotlib() -> None:
    import matplotlib.pyplot as plt

    plt.rcParams["font.sans-serif"] = [
        "Noto Sans CJK SC",
        "Microsoft YaHei",
        "SimHei",
        "PingFang SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def save_figure(figure: Any, path: Path) -> None:
    figure.tight_layout()
    figure.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    import matplotlib.pyplot as plt

    plt.close(figure)


def draw_charts(
    output_dir: Path,
    detail: pd.DataFrame,
    projects: pd.DataFrame,
    confidence_bins: pd.DataFrame,
    strategy_comparison: pd.DataFrame,
) -> list[Path]:
    import matplotlib.pyplot as plt
    from matplotlib.ticker import PercentFormatter

    configure_matplotlib()
    paths: list[Path] = []

    path = output_dir / "01_三种策略覆盖率准确率对比.png"
    labels = strategy_comparison["策略"].tolist()
    score_thresholds = strategy_comparison["参考置信分阈值"].dropna()
    score_threshold = float(score_thresholds.iloc[0]) if not score_thresholds.empty else math.nan
    threshold_text = f"模型分阈值={score_threshold:.4f}" if not math.isnan(score_threshold) else "模型分阈值未设置"
    x = np.arange(len(labels))
    width = 0.34
    fig, ax = plt.subplots(figsize=(10, 6))
    coverage_bars = ax.bar(
        x - width / 2,
        strategy_comparison["自动放行覆盖率"],
        width,
        color="#4c78a8",
        label="自动放行覆盖率",
    )
    accuracy_bars = ax.bar(
        x + width / 2,
        strategy_comparison["自动放行准确率"],
        width,
        color="#2a9d8f",
        label="自动放行准确率",
    )
    ax.set_xticks(x, labels)
    ax.set_ylim(0, 1.12)
    ax.yaxis.set_major_formatter(PercentFormatter(1.0))
    ax.set_ylabel("比例")
    ax.set_title(f"三种自动放行策略直接对比（{threshold_text}）\n覆盖率越高处理量越大，准确率越高放错越少")
    ax.legend(loc="lower right")
    ax.grid(axis="y", alpha=0.2)
    for bars in (coverage_bars, accuracy_bars):
        for bar in bars:
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height() + 0.015,
                f"{bar.get_height():.2%}",
                ha="center",
                va="bottom",
                fontsize=10,
            )
    save_figure(fig, path)
    paths.append(path)

    path = output_dir / "02_三种策略错误控制对比.png"
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5.8))
    colors = ["#e76f51", "#f4a261", "#2a9d8f"]
    error_bars = ax1.bar(labels, strategy_comparison["自动放行错误数"], color=colors)
    ax1.set_ylabel("错误条数")
    ax1.set_title(f"自动放行中的错误数（{threshold_text}）\n越少越安全")
    ax1.grid(axis="y", alpha=0.2)
    for bar, value in zip(error_bars, strategy_comparison["自动放行错误数"]):
        ax1.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{int(value):,}", ha="center", va="bottom")
    intercept_bars = ax2.bar(labels, strategy_comparison["错误拦截率"], color=colors)
    ax2.set_ylim(0, 1.08)
    ax2.yaxis.set_major_formatter(PercentFormatter(1.0))
    ax2.set_ylabel("错误拦截率")
    ax2.set_title("进入审核的错误占全部错误比例\n越高越安全")
    ax2.grid(axis="y", alpha=0.2)
    for bar, value in zip(intercept_bars, strategy_comparison["错误拦截率"]):
        ax2.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.015, f"{value:.2%}", ha="center", va="bottom")
    save_figure(fig, path)
    paths.append(path)

    path = output_dir / "03_置信分区间准确率.png"
    confidence_plot = confidence_bins.loc[confidence_bins["样本数"].gt(0)].copy()
    x = np.arange(len(confidence_plot))
    fig, ax1 = plt.subplots(figsize=(11, 6))
    count_bars = ax1.bar(x, confidence_plot["样本数"], color="#90caf9", alpha=0.85, label="样本数")
    ax1.set_xticks(x, confidence_plot["置信分区间"], rotation=25, ha="right")
    ax1.set_xlabel("模型置信分区间")
    ax1.set_ylabel("样本数")
    ax2 = ax1.twinx()
    ax2.plot(x, confidence_plot["准确率"], color="#d1495b", marker="o", linewidth=2.4, label="真实准确率")
    ax2.set_ylim(max(0, float(confidence_plot["准确率"].min()) - 0.08), 1.02)
    ax2.yaxis.set_major_formatter(PercentFormatter(1.0))
    ax2.set_ylabel("真实准确率")
    ax1.set_title("模型置信分能否反映真实正确率\n柱形是样本数，红线是该分数区间的真实准确率")
    for bar, value in zip(count_bars, confidence_plot["样本数"]):
        ax1.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"{int(value):,}", ha="center", va="bottom", fontsize=9)
    for index, value in enumerate(confidence_plot["准确率"]):
        ax2.text(index, min(1.01, value + 0.012), f"{value:.1%}", ha="center", va="bottom", color="#a4133c", fontsize=9)
    save_figure(fig, path)
    paths.append(path)

    path = output_dir / "04_项目分流异常散点.png"
    fig, ax = plt.subplots(figsize=(10, 6.5))
    flagged = projects["是否重点异常"].astype(bool)
    sizes = 35 + 260 * np.sqrt(projects["样本数"] / projects["样本数"].max())
    ax.scatter(projects.loc[~flagged, "简单占比"], projects.loc[~flagged, "简单准确率"], s=sizes[~flagged], alpha=0.55, color="#2a9d8f", label="正常/观察")
    ax.scatter(projects.loc[flagged, "简单占比"], projects.loc[flagged, "简单准确率"], s=sizes[flagged], alpha=0.8, color="#e63946", label="重点异常")
    project_labels = projects.loc[flagged].sort_values(["简单准确率", "错误数"], ascending=[True, False]).head(8)
    offsets = ((6, 8), (-6, 12), (8, -16), (-8, -18))
    for position, (_, row) in enumerate(project_labels.iterrows()):
        ax.annotate(
            str(row["项目名称"]),
            (row["简单占比"], row["简单准确率"]),
            fontsize=8,
            xytext=offsets[position % len(offsets)],
            textcoords="offset points",
            ha="left" if position % 2 == 0 else "right",
            bbox={"boxstyle": "round,pad=0.15", "facecolor": "white", "alpha": 0.7, "edgecolor": "none"},
        )
    ax.axvline(float(detail["分析_分流类别"].eq("简单").mean()), color="#777", linestyle="--", linewidth=1)
    ax.axhline(float(detail.loc[detail["分析_分流类别"].eq("简单"), "分析_编码是否正确"].mean()), color="#777", linestyle="--", linewidth=1)
    ax.set_xlim(-0.02, 1.02)
    ax.set_ylim(0, 1.02)
    ax.set_xlabel("项目简单占比")
    ax.set_ylabel("项目简单样本准确率")
    ax.set_title("项目分流异常定位\n越靠右下越危险：判为简单的比例高，但简单样本准确率低")
    ax.text(0.98, 0.04, "重点关注：放行多、准确率低", transform=ax.transAxes, ha="right", color="#b00020", fontsize=10)
    ax.legend()
    save_figure(fig, path)
    paths.append(path)
    return paths


def _format_excel(output_path: Path, sheet_names: list[str]) -> None:
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="1F5B66")
    anomaly_fill = PatternFill("solid", fgColor="FDE2E4")
    for sheet_name in sheet_names:
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        if sheet.max_row >= 1 and sheet.max_column >= 1:
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        headers = {cell.column: str(cell.value or "") for cell in sheet[1]}
        for column_index, header in headers.items():
            letter = sheet.cell(1, column_index).column_letter
            sample_values = [str(sheet.cell(row, column_index).value or "") for row in range(1, min(sheet.max_row, 300) + 1)]
            width = min(60, max(10, max(map(len, sample_values), default=10) + 2))
            sheet.column_dimensions[letter].width = width
            if any(keyword in header for keyword in ("准确率", "错误率", "占比", "覆盖率", "拦截率", "校准差")):
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, column_index).number_format = "0.00%"
                if sheet.max_row >= 2:
                    sheet.conditional_formatting.add(
                        f"{letter}2:{letter}{sheet.max_row}",
                        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
                    )
            elif "置信分" in header or header in {"值"}:
                for row in range(2, sheet.max_row + 1):
                    if isinstance(sheet.cell(row, column_index).value, float):
                        sheet.cell(row, column_index).number_format = "0.0000"
        if sheet_name == "项目分析":
            anomaly_column = next((column for column, header in headers.items() if header == "是否重点异常"), None)
            if anomaly_column:
                for row in range(2, sheet.max_row + 1):
                    if sheet.cell(row, anomaly_column).value:
                        for cell in sheet[row]:
                            cell.fill = anomaly_fill
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
    workbook.save(output_path)


def export_excel(
    output_path: Path,
    tables: dict[str, pd.DataFrame],
    runtime_rows: list[dict[str, Any]],
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in tables.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        pd.DataFrame(runtime_rows).to_excel(writer, sheet_name="运行配置", index=False)
    _format_excel(output_path, [*tables, "运行配置"])


def export_text_report(
    report_path: Path,
    conclusions: pd.DataFrame,
    strategy_comparison: pd.DataFrame,
) -> None:
    lines = ["分流与模型置信分分析结论", "=" * 32, ""]
    for _, row in conclusions.iterrows():
        lines.append(f"{int(row['序号'])}. {row['主题']}：{row['结论']}")
    lines.extend(["", "三种策略关键指标", "-" * 32])
    for _, row in strategy_comparison.iterrows():
        lines.append(
            f"{row['策略']}：覆盖率 {row['自动放行覆盖率']:.2%}，"
            f"放行准确率 {row['自动放行准确率']:.2%}，"
            f"放行错误 {int(row['自动放行错误数']):,} 条，"
            f"错误拦截率 {row['错误拦截率']:.2%}。"
        )
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def ensure_outputs(config: AnalysisConfig, output_path: Path, image_paths: list[Path] | None = None) -> None:
    existing = [path for path in [output_path, *(image_paths or [])] if path.exists()]
    if existing and not config.overwrite:
        formatted = "\n".join(f"- {path}" for path in existing)
        raise FileExistsError(f"输出文件已存在，请更换目录或增加 --overwrite:\n{formatted}")
    config.output_dir.mkdir(parents=True, exist_ok=True)


def run_analysis(config: AnalysisConfig) -> dict[str, Any]:
    output_path = config.output_dir / "分流置信分分析.xlsx"
    report_path = config.output_dir / "分析结论.txt"
    expected_chart_paths = [
        config.output_dir / "01_三种策略覆盖率准确率对比.png",
        config.output_dir / "02_三种策略错误控制对比.png",
        config.output_dir / "03_置信分区间准确率.png",
        config.output_dir / "04_项目分流异常散点.png",
    ]
    ensure_outputs(config, output_path, [report_path, *expected_chart_paths])
    LOGGER.info("读取 Excel: %s", config.input_path)
    detail, columns, quality = load_and_prepare(config)
    LOGGER.info("有效数据 %d 条，项目 %d 个", len(detail), detail["分析_项目名称"].nunique())

    summary, routing = build_routing_stats(detail)
    projects = build_project_stats(detail, config.min_project_samples)
    confidence_bins = build_confidence_bins(detail)
    calibration, confidence_metrics = build_calibration(detail)
    simulation = build_threshold_simulation(detail)
    recommendations = build_recommendations(simulation, config.target_accuracies)
    quadrants = build_quadrants(detail, config.reference_threshold)
    reason_combinations, atomic_reasons = build_reason_stats(detail)
    fields = build_field_stats(detail)
    review_cross = build_review_cross(detail)
    strategy_comparison = build_strategy_comparison(detail, config.reference_threshold)
    conclusions = build_conclusions(summary, strategy_comparison, recommendations, confidence_metrics)
    image_guide = build_image_guide(config.reference_threshold)

    chart_paths = draw_charts(
        config.output_dir,
        detail,
        projects,
        confidence_bins,
        strategy_comparison,
    )
    key_columns = [
        column for column in (
            "主表id",
            "子表id",
            "材料描述",
            "材料描述(多行)",
            columns["project"],
            "分类",
            TRUTH_COLUMN,
            PREDICTED_COLUMN,
            columns["confidence"],
            columns["difficulty"],
            REASON_COLUMN,
            REVIEW_COLUMN,
            "分析_Excel行号",
            "分析_模型置信分",
            "分析_分流类别",
            "分析_编码是否正确",
            "分析_四象限",
        ) if column in detail.columns
    ]
    analysis_detail = detail[key_columns].copy()
    error_detail = analysis_detail.loc[~detail["分析_编码是否正确"]].copy()
    tables = {
        "结论摘要": conclusions,
        "三种策略对比": strategy_comparison,
        "图片说明": image_guide,
        "分析摘要": summary,
        "数据质量": quality,
        "分流统计": routing,
        "项目分析": projects,
        "异常项目": projects.loc[projects["是否重点异常"]].copy(),
        "置信分分箱": confidence_bins,
        "置信分校准": calibration,
        "置信分指标": confidence_metrics,
        "阈值策略模拟": simulation,
        "阈值推荐": recommendations,
        "四象限分析": quadrants,
        "分流原因组合": reason_combinations,
        "分流原因规则": atomic_reasons,
        "字段错误": fields,
        "审核标记交叉": review_cross,
        "错误明细": error_detail,
        "分析明细": analysis_detail,
    }
    runtime_rows = [
        {"配置项": "输入Excel", "值": str(config.input_path)},
        {"配置项": "输出目录", "值": str(config.output_dir)},
        {"配置项": "工作表", "值": str(config.sheet)},
        {"配置项": "项目列", "值": columns["project"]},
        {"配置项": "置信分列", "值": columns["confidence"]},
        {"配置项": "分流列", "值": columns["difficulty"]},
        {"配置项": "参考置信分阈值", "值": config.reference_threshold},
        {"配置项": "项目异常最少样本", "值": config.min_project_samples},
        {"配置项": "自动放行目标准确率", "值": ", ".join(f"{value:.3%}" for value in config.target_accuracies)},
        {"配置项": "正确性口径", "值": f"{PREDICTED_COLUMN} 与 {TRUTH_COLUMN} 经NFKC、去空白、转大写后完全一致"},
    ]
    LOGGER.info("写入分析 Excel: %s", output_path)
    export_excel(output_path, tables, runtime_rows)
    export_text_report(report_path, conclusions, strategy_comparison)
    return {
        "excel": output_path,
        "report": report_path,
        "images": chart_paths,
        "summary": summary,
        "recommendations": recommendations,
        "strategy_comparison": strategy_comparison,
    }


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(asctime)s [%(levelname)s] %(message)s")
    config = build_config(args)
    result = run_analysis(config)
    print(result["summary"].to_string(index=False))
    print("\n阈值推荐:")
    print(result["recommendations"].to_string(index=False))
    print("\n三种策略对比:")
    print(result["strategy_comparison"].to_string(index=False))
    print(f"\n分析 Excel: {result['excel']}")
    print(f"分析结论: {result['report']}")
    for path in result["images"]:
        print(f"分析图片: {path}")


if __name__ == "__main__":
    main()
