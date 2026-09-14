import csv
import io
import json
import tempfile
import unittest
from pathlib import Path

from apps.platform.batch_export import iter_csv, iter_stage1_json, write_xlsx


class BatchExportTest(unittest.TestCase):
    def setUp(self) -> None:
        self.items = [
            {"index": 0, "text": "原始描述A", "project_name": "项目A", "category": "直管"},
            {"index": 1, "text": "原始描述B", "project_name": "项目B", "category": "管件"},
        ]
        self.result = {
            "original_text": "原始描述A",
            "processed_text": "格式化描述A",
            "final_code": "P50",
            "success": True,
            "need_review": False,
            "confidence": 0.99,
            "fields": {
                "TYPE": {"stage1_raw": {"value": {"BODY": "直管"}}, "stage2_output": {"code": "P"}},
                "SIZE": {
                    "stage1_raw": {
                        "value": {
                            "ITEMS": [{"SCOPE": "BODY", "ROLE": "SINGLE", "SIZE": [{"type": "DN", "value": "50"}], "THICKNESS": []}],
                            "LENGTH": "",
                            "PRESSURE": "",
                        }
                    },
                    "stage2_output": {"code": "50"},
                },
            },
        }
        self.rows = [(0, 0, self.result)]

    def test_stage1_export_contains_all_processed_rows_and_flattens_structure(self) -> None:
        payload = json.loads(b"".join(iter_stage1_json(self.items, self.rows)))

        self.assertEqual(len(payload), 1)
        self.assertEqual(payload[0]["original_input"], "原始描述A")
        self.assertEqual(payload[0]["input"], "格式化描述A")
        self.assertIn("ITEMS", payload[0]["output"])
        self.assertNotIn("SIZE", payload[0]["output"])

    def test_csv_export_keeps_unprocessed_items_as_blank_rows(self) -> None:
        text = b"".join(iter_csv(self.items, self.rows)).decode("utf-8-sig")
        records = list(csv.DictReader(io.StringIO(text)))

        self.assertEqual(len(records), 2)
        self.assertEqual(records[0]["原始总编码"], "P50")
        self.assertEqual(records[1]["原始描述"], "原始描述B")
        self.assertEqual(records[1]["原始总编码"], "")

    def test_xlsx_export_keeps_all_items(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "results.xlsx"
            write_xlsx(path, self.items, self.rows)
            workbook = load_workbook(path, read_only=True)
            sheet = workbook["编码结果"]
            values = list(sheet.iter_rows(values_only=True))
            workbook.close()

        self.assertEqual(len(values), 3)
        self.assertEqual(values[1][3], "原始描述A")
        self.assertEqual(values[2][3], "原始描述B")


if __name__ == "__main__":
    unittest.main()
